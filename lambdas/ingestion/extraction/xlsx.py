import io
import openpyxl

class XLSXExtractor:
    def extract(self, file_bytes: bytes) -> list[dict]:
        """
        Extracts text from XLSX files using openpyxl.
        Preserves sheet names.
        Returns a list of dicts with 'page_number' and 'text'.
        We map 'page_number' to the sheet index.
        """
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        pages = []
        
        for i, sheet_name in enumerate(wb.sheetnames, start=1):
            sheet = wb[sheet_name]
            rows_text = []
            
            # Simple row extraction
            for row in sheet.iter_rows(values_only=True):
                # Filter out None values and convert to string
                row_values = [str(cell) for cell in row if cell is not None and str(cell).strip() != ""]
                if row_values:
                    rows_text.append(" | ".join(row_values))
                    
            if rows_text:
                sheet_text = f"Sheet: {sheet_name}\n\n" + "\n".join(rows_text)
                pages.append({
                    "page_number": i,
                    "text": sheet_text
                })
                
        return pages
