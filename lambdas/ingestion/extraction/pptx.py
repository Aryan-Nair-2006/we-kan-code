import io
from pptx import Presentation

class PPTXExtractor:
    def extract(self, file_bytes: bytes) -> list[dict]:
        """
        Extracts text from PPTX files using python-pptx.
        Preserves slide numbers.
        Returns a list of dicts with 'page_number' and 'text'.
        """
        prs = Presentation(io.BytesIO(file_bytes))
        pages = []
        
        for i, slide in enumerate(prs.slides, start=1):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text.strip())
                elif shape.has_table:
                    for row in shape.table.rows:
                        row_vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_vals:
                            slide_text.append(" | ".join(row_vals))
                            
            if slide_text:
                pages.append({
                    "page_number": i,
                    "text": "\n".join(slide_text)
                })
                
        return pages
